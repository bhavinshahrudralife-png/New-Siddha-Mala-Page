from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, File, UploadFile
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Dict, Any
from datetime import datetime, timezone, timedelta
import os
import uuid
import logging
import bcrypt
import jwt
from login_limits import account_key, check_lock, record_failure
from enquiry_models import EnquiryCreate, Enquiry, LeadUpdate, LeadPage
from page_content import PageSave, PageResponse, read_page, save_page, seed_page
from media_storage import MediaItem, upload_media, serve_media, init_storage

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

JWT_ALGORITHM = "HS256"

def get_jwt_secret() -> str:
    return os.environ["JWT_SECRET"]

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))

def create_access_token(user_id: str, email: str) -> str:
    payload = {"sub": user_id, "email": email, "type": "access",
               "exp": datetime.now(timezone.utc) + timedelta(hours=12)}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(status_code=401, detail="Invalid token type")
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
    if not user:
        raise HTTPException(status_code=401, detail="User not found")
    return user

# ---------- Models ----------
class LoginInput(BaseModel):
    email: str
    password: str

DEFAULT_CONTENT: Dict[str, Any] = {
    "hero_title": "One mala. A complete spiritual companion.",
    "hero_subtitle": "Handcrafted with authentic Rudraksha. A sacred combination of 1 to 14 Mukhi, Gauri Shankar and Ganesh.",
    "phone": "+91 72088 19922",
    "email": "info@rudralife.com",
    "whatsapp": "919619990717",
    "address": "901, Sri Krishna, New Link Road, Andheri (West), Mumbai — 400 053",
    "hours": "Mon – Sat · 10:00 AM to 7:00 PM (IST)",
    "website": "www.rudralife.com",
    "stat_certified": "100%",
    "stat_years": "25+",
    "stat_souls": "5,00,000+",
}

# ---------- Public ----------
@api_router.get("/")
async def root():
    return {"message": "Rudralife Siddha Mala API"}

@api_router.post("/enquiries", response_model=Enquiry)
async def create_enquiry(input: EnquiryCreate):
    from pymongo.errors import DuplicateKeyError
    existing = await db.enquiries.find_one({'submission_id': input.submission_id}, {'_id': 0})
    if existing: return Enquiry(**existing)
    settings = await db.page_content.find_one({'key': 'landing'}, {'_id': 0, 'page.form.fields': 1})
    fields = (settings or {}).get('page', {}).get('form', {}).get('fields', [])
    for field in fields:
        if field.get('required') and field.get('key') in ('city', 'company', 'profession', 'message') and not getattr(input, field['key']).strip():
            raise HTTPException(422, f"{field.get('label', field['key'])} is required.")
    enquiry = Enquiry(**input.model_dump())
    try:
        await db.enquiries.insert_one({**enquiry.model_dump(), 'submission_id': input.submission_id})
    except DuplicateKeyError:
        stored = await db.enquiries.find_one({'submission_id': input.submission_id}, {'_id': 0})
        return Enquiry(**stored)
    return enquiry

@api_router.get('/page-content', response_model=PageResponse)
async def public_page():
    return await read_page(db)

@api_router.get('/media/{media_id}')
async def public_media(media_id: str):
    return await serve_media(db, media_id)

@api_router.get("/content")
async def get_content():
    doc = await db.site_content.find_one({"key": "main"}, {"_id": 0})
    if not doc:
        return DEFAULT_CONTENT
    doc.pop("key", None)
    return doc

# ---------- Auth ----------
@api_router.post("/auth/login")
async def login(input: LoginInput, response: Response, request: Request):
    email = input.email.lower().strip()
    identifier = account_key(email)
    await check_lock(db, identifier)
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(input.password, user["password_hash"]):
        await record_failure(db, identifier)
        raise HTTPException(status_code=401, detail="Invalid email or password")
    await db.login_attempts.delete_one({"_id": identifier})
    token = create_access_token(user["id"], email)
    response.set_cookie(key="access_token", value=token, httponly=True, secure=True,
                        samesite="none", max_age=43200, path="/")
    return {"id": user["id"], "email": email, "name": user.get("name", "Admin"), "role": user.get("role", "admin")}

@api_router.get("/auth/me")
async def me(user=Depends(get_current_user)):
    return user

@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"message": "Logged out"}

# ---------- Admin ----------
def lead_filter(query='', status=''):
    import re
    filters = {}
    if query.strip():
        pattern = {'$regex': re.escape(query.strip()[:150]), '$options': 'i'}
        filters['$or'] = [{field: pattern} for field in ('name', 'email', 'phone', 'city', 'company', 'profession')]
    if status:
        if status not in ('new', 'contacted', 'qualified', 'closed'):
            raise HTTPException(422, 'Invalid lead status')
        filters['status'] = {'$in': ['new', None]} if status == 'new' else status
    return filters

@api_router.get('/admin/leads/page', response_model=LeadPage)
async def paged_leads(query: str = '', status: str = '', page: int = 1, user=Depends(get_current_user)):
    if page < 1: raise HTTPException(422, 'Invalid page')
    filters = lead_filter(query, status)
    total = await db.enquiries.count_documents(filters)
    docs = await db.enquiries.find(filters, {'_id': 0}).sort('created_at', -1).skip((page-1)*25).limit(25).to_list(25)
    return LeadPage(items=[Enquiry(**d) for d in docs], total=total, page=page, has_more=page*25 < total)

@api_router.get('/admin/leads/export')
async def lead_export(query: str = '', status: str = '', user=Depends(get_current_user)):
    import csv, io
    from starlette.responses import StreamingResponse
    keys = ['name', 'email', 'country', 'country_code', 'phone', 'company', 'profession', 'city', 'message', 'interest', 'source', 'status', 'notes', 'created_at']
    async def rows():
        yield '\ufeff'
        buffer = io.StringIO(); writer = csv.writer(buffer); writer.writerow(keys); yield buffer.getvalue()
        async for lead in db.enquiries.find(lead_filter(query,status), {'_id':0}).sort('created_at',-1):
            buffer.seek(0); buffer.truncate()
            values=[]
            for key in keys:
                value=str(lead.get(key, ''))
                values.append("'"+value if value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else value)
            writer.writerow(values); yield buffer.getvalue()
    return StreamingResponse(rows(), media_type='text/csv', headers={'Content-Disposition':'attachment; filename="rudralife-enquiries.csv"'})

LEAD_COLUMNS = [('Received', 'created_at'), ('Full name', 'name'), ('Email', 'email'), ('Country', 'country'), ('Country code', 'country_code'), ('Phone', 'phone'), ('Company', 'company'), ('Profession', 'profession'), ('City', 'city'), ('Message', 'message'), ('Interest', 'interest'), ('Form', 'source'), ('Status', 'status'), ('Notes', 'notes')]

@api_router.get('/admin/leads/export.xlsx')
async def lead_export_excel(query: str = '', status: str = '', user=Depends(get_current_user)):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from starlette.responses import Response as RawResponse
    wb = Workbook(); ws = wb.active; ws.title = 'Enquiries'
    ws.append([title for title, _ in LEAD_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='642C2B')
    async for lead in db.enquiries.find(lead_filter(query, status), {'_id': 0}).sort('created_at', -1):
        row = []
        for _, key in LEAD_COLUMNS:
            value = str(lead.get(key) or '')
            if key == 'created_at' and value:
                value = value[:19].replace('T', ' ')
            row.append("'" + value if value.startswith(('=', '+', '-', '@')) and key not in ('phone', 'country_code') else value)
        ws.append(row)
    for column, width in zip('ABCDEFGHIJKLMN', (20, 24, 30, 9, 12, 18, 24, 20, 16, 50, 24, 10, 12, 30)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = 'A2'
    buffer = io.BytesIO(); wb.save(buffer)
    return RawResponse(buffer.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="rudralife-enquiries.xlsx"'})

@api_router.put('/admin/page-content', response_model=PageResponse)
async def edit_page(data: PageSave, user=Depends(get_current_user)):
    return await save_page(db, data)

@api_router.post('/admin/media', response_model=MediaItem)
async def media_upload(file: UploadFile = File(...), user=Depends(get_current_user)):
    return await upload_media(db, file, user['id'])

@api_router.get('/admin/media', response_model=List[MediaItem])
async def media_list(user=Depends(get_current_user)):
    return await db.media.find({'is_deleted': False}, {'_id': 0, 'storage_path': 0, 'gridfs_id': 0}).sort('created_at', -1).to_list(500)

@api_router.delete('/admin/media/{media_id}')
async def media_delete(media_id: str, user=Depends(get_current_user)):
    result = await db.media.update_one({'id': media_id, 'is_deleted': False}, {'$set': {'is_deleted': True}})
    if not result.matched_count: raise HTTPException(404, 'Media not found.')
    return {'deleted': True}

@api_router.patch('/admin/leads/{lead_id}', response_model=Enquiry)
async def edit_lead(lead_id: str, data: LeadUpdate, user=Depends(get_current_user)):
    from pymongo import ReturnDocument
    lead = await db.enquiries.find_one_and_update({'id': lead_id}, {'$set': data.model_dump()},
        return_document=ReturnDocument.AFTER, projection={'_id': 0})
    if not lead: raise HTTPException(404, 'Lead not found')
    return Enquiry(**lead)

@api_router.get("/admin/leads", response_model=List[Enquiry])
async def get_leads(user=Depends(get_current_user)):
    leads = await db.enquiries.find({}, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return leads

@api_router.delete("/admin/leads/{lead_id}")
async def delete_lead(lead_id: str, user=Depends(get_current_user)):
    result = await db.enquiries.delete_one({"id": lead_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Lead not found")
    return {"message": "Deleted"}

@api_router.put("/admin/content")
async def update_content(content: Dict[str, Any], user=Depends(get_current_user)):
    content.pop("_id", None)
    content.pop("key", None)
    await db.site_content.update_one({"key": "main"}, {"$set": {**content, "key": "main"}}, upsert=True)
    mapping = {'hero_title':'page.hero.tagline', 'hero_subtitle':'page.hero.description',
        **{key:f'page.footer.{key}' for key in ('phone','email','whatsapp','address','hours','website')}}
    changes = {mapping[key]:value for key,value in content.items() if key in mapping}
    if changes:
        await db.page_content.update_one({'key':'landing'}, {'$set':changes, '$inc':{'revision':1}})
    return {"message": "Content updated"}

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=[os.environ["FRONTEND_URL"]],
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

@app.on_event("startup")
async def startup():
    await db.page_content.create_index('key', unique=True)
    await db.enquiries.create_index('submission_id', unique=True, sparse=True)
    await db.enquiries.create_index('created_at')
    await db.media.create_index('id', unique=True)
    try:
        await init_storage()
        logger.info('Media storage initialized')
    except Exception:
        logger.warning('Media storage initialization deferred until next upload')
    await db.users.create_index("email", unique=True)
    await db.login_attempts.create_index("expires_at", expireAfterSeconds=0)
    admin_email = os.environ["ADMIN_EMAIL"].lower()
    admin_password = os.environ["ADMIN_PASSWORD"]
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({
            "id": str(uuid.uuid4()), "email": admin_email,
            "password_hash": hash_password(admin_password),
            "name": "Rudralife Admin", "role": "admin",
            "created_at": datetime.now(timezone.utc).isoformat()})
        logger.info("Admin user seeded")
    elif not verify_password(admin_password, existing["password_hash"]):
        await db.users.update_one({"email": admin_email},
                                  {"$set": {"password_hash": hash_password(admin_password)}})
    if not await db.site_content.find_one({"key": "main"}):
        await db.site_content.insert_one({**DEFAULT_CONTENT, "key": "main"})
    # Replace only the previous starter copy; never overwrite admin custom edits.
    legacy_copy = {
        "hero_title": "Are you searching for something more — a stillness, a strength, a sense of destiny awakening within you?",
        "hero_subtitle": "For thousands of years, seekers, sages, and kings turned to a single sacred treasure to align mind, body and soul — the Siddha Mala. Today, that same sacred power finds its way to those who are ready.",
        "phone": "+91 98203 02028",
        "whatsapp": "919820302028",
        "address": "Rudralife®, Mumbai, Maharashtra, India",
    }
    for field, previous in legacy_copy.items():
        await db.site_content.update_one(
            {"key": "main", field: previous}, {"$set": {field: DEFAULT_CONTENT[field]}}
        )
    legacy = await db.site_content.find_one({'key':'main'}, {'_id':0})
    await seed_page(db, legacy or DEFAULT_CONTENT)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
