# Iteration 11: Fifth edition – company field, xlsx export, CMS sections, media upload (GridFS fallback)
import io
import os
import uuid
import pytest
import requests
from openpyxl import load_workbook
from PIL import Image
from pymongo import MongoClient

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', '').rstrip('/')
if not BASE_URL:
    with open('/app/frontend/.env') as f:
        BASE_URL = [l.split('=', 1)[1].strip() for l in f if l.startswith('REACT_APP_BACKEND_URL')][0].rstrip('/')


def creds():
    text = open('/app/memory/test_credentials.md').read()
    return 'admin@rudralife.com', 'Rudralife@2025' if 'Rudralife@2025' in text else None


@pytest.fixture(scope='module')
def admin():
    s = requests.Session()
    email, pw = creds()
    r = s.post(f'{BASE_URL}/api/auth/login', json={'email': email, 'password': pw})
    if r.status_code != 200:
        pytest.skip('admin login failed')
    return s


@pytest.fixture(scope='module')
def qa_lead(admin):
    payload = {'name': 'TEST_QA Edition5', 'email': f'test_{uuid.uuid4().hex[:6]}@example.com',
               'phone': '9876543210', 'country': 'IN', 'country_code': '+91', 'company': 'TEST_Company Ltd',
               'profession': 'TEST_Engineer', 'message': 'TEST_message hello', 'source': 'popup',
               'submission_id': str(uuid.uuid4())}
    r = requests.post(f'{BASE_URL}/api/enquiries', json=payload)
    assert r.status_code == 200, r.text
    data = r.json()
    yield data
    admin.delete(f"{BASE_URL}/api/admin/leads/{data['id']}")


# Public page content structure
def test_page_order_and_removed_sections():
    p = requests.get(f'{BASE_URL}/api/page-content').json()['page']
    order = p['order']
    for removed in ('intro', 'beadJourney', 'education', 'community', 'experts', 'consultation'):
        assert removed not in order
    assert order.index('gallery') < order.index('story') < order.index('beads')
    assert order[-1] == 'featured' and 'founder' in order
    assert len(p['beads']['items']) == 16
    assert [x['id'] for x in p['patrons']['items']] == ['milind', 'dimple', 'archana', 'sunil']
    assert p['patrons']['interval_seconds'] == 5
    assert p['founder']['interval_seconds'] == 2 and len(p['founder']['images']) == 3
    assert len(p['featured']['items']) == 9
    assert [f['key'] for f in p['form']['fields']] == ['name', 'email', 'phone', 'company', 'profession', 'message']


# Enquiry company persisted
def test_enquiry_company_persisted(admin, qa_lead):
    assert qa_lead['company'] == 'TEST_Company Ltd'
    r = admin.get(f'{BASE_URL}/api/admin/leads/page', params={'query': 'TEST_Company'})
    assert r.status_code == 200
    items = [i for i in r.json()['items'] if i['id'] == qa_lead['id']]
    assert items and items[0]['profession'] == 'TEST_Engineer' and items[0]['phone'] == '+919876543210'


# Excel export
def test_xlsx_export(admin, qa_lead):
    r = admin.get(f'{BASE_URL}/api/admin/leads/export.xlsx')
    assert r.status_code == 200
    assert 'spreadsheetml' in r.headers['content-type']
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    for h in ('Company', 'Profession', 'Message', 'Full name', 'Phone'):
        assert h in headers
    names = [row[headers.index('Full name')] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert 'TEST_QA Edition5' in names
    row = [r_ for r_ in ws.iter_rows(min_row=2, values_only=True) if r_[headers.index('Full name')] == 'TEST_QA Edition5'][0]
    assert row[headers.index('Company')] == 'TEST_Company Ltd'
    assert row[headers.index('Phone')] == '+919876543210'


def test_xlsx_export_requires_auth():
    r = requests.get(f'{BASE_URL}/api/admin/leads/export.xlsx')
    assert r.status_code == 401


def test_csv_export_has_company(admin, qa_lead):
    r = admin.get(f'{BASE_URL}/api/admin/leads/export')
    assert r.status_code == 200
    text = r.content.decode('utf-8-sig')
    assert 'company' in text.splitlines()[0] and 'TEST_Company Ltd' in text


# Media upload (PNG + JPG) and serve, then mark not used
@pytest.mark.parametrize('fmt,ctype', [('PNG', 'image/png'), ('JPEG', 'image/jpeg')])
def test_media_upload_and_serve(admin, fmt, ctype):
    buf = io.BytesIO()
    mode = 'RGBA' if fmt == 'PNG' else 'RGB'
    Image.new(mode, (40, 40), (200, 100, 50) if mode == 'RGB' else (200, 100, 50, 128)).save(buf, format=fmt)
    r = admin.post(f'{BASE_URL}/api/admin/media', files={'file': (f'TEST_qa.{fmt.lower()}', buf.getvalue(), ctype)})
    assert r.status_code == 200, r.text
    item = r.json()
    g = requests.get(f"{BASE_URL}{item['url']}")
    assert g.status_code == 200 and g.headers['content-type'].startswith('image/')
    Image.open(io.BytesIO(g.content)).verify()
    # cleanup: mark QA media not used (no delete API exists)
    env = dict(l.strip().split('=', 1) for l in open('/app/backend/.env') if '=' in l)
    mc = MongoClient(env['MONGO_URL'].strip('"'))
    mc[env['DB_NAME'].strip('"')].media.update_one({'id': item['id']}, {'$set': {'is_deleted': True}})


def test_media_rejects_svg(admin):
    r = admin.post(f'{BASE_URL}/api/admin/media', files={'file': ('x.svg', b'<svg xmlns="http://www.w3.org/2000/svg"/>', 'image/svg+xml')})
    assert r.status_code == 422
